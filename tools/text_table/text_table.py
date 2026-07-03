# -*- coding: utf-8 -*-
"""
把游戏里的所有文本汇总到一个多页签 Excel 里配置，改完再一键转回 JSON。

用法（在项目根目录运行）:
    # JSON  ->  Excel（把当前所有文本导出成表格，用来初始化 / 刷新表格）
    python tools/text_table/text_table.py export

    # Excel ->  JSON（把表格里改好的文本写回游戏用的 JSON）
    python tools/text_table/text_table.py import

可选参数: --xlsx <路径>   自定义 Excel 文件位置（默认 tools/text_table/game_texts.xlsx）

页签说明:
    story        普通线剧本（story.json / story_en.json）
    story_true   真结局剧本（story_true.json / story_true_en.json）
    pet          普通桌宠对白（pet_dialogues.json / _en.json）
    pet_true     真结局桌宠对白（pet_dialogues_true.json / _en.json）
    ending_avg   真结局 AVG 悬浮框（ending_avg.json / _en.json）
    meta         起始节点、彩蛋、历史屏蔽语、AVG/桌宠里少量固定文本

导入时按 id / 分区+序号 合并，不在表格里的结构字段（setVariables、
conditionalNexts、桌宠计时参数、animations 帧动画等）会从旧 JSON 原样保留。
日语配音文本（text_ja 列）会同步写回 tools/tts/voice_lines_ja.json。
"""
import argparse
import json
import shutil
import sys
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("缺少 openpyxl，请先运行:  python -m pip install openpyxl")
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[2]
SA = ROOT / "Assets" / "StreamingAssets"
JA_PATH = ROOT / "tools" / "tts" / "voice_lines_ja.json"
DEFAULT_XLSX = ROOT / "tools" / "text_table" / "game_texts.xlsx"
BACKUP_DIR = ROOT / "tools" / "text_table" / "backups"
_BACKUP_STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")

STORY_FILES = {
    "story": ("story.json", "story_en.json"),
    "story_true": ("story_true.json", "story_true_en.json"),
}
PET_FILES = {
    "pet": ("pet_dialogues.json", "pet_dialogues_en.json"),
    "pet_true": ("pet_dialogues_true.json", "pet_dialogues_true_en.json"),
}
AVG_FILES = ("ending_avg.json", "ending_avg_en.json")

PET_SECTIONS = ["idle", "click", "morning", "noon", "evening", "night", "welcome", "headpat"]
AVG_SECTIONS = ["enter", "lines"]
MAX_CHOICES = 4

STORY_HEADERS = [
    "id", "speaker_zh", "speaker_en", "text_zh", "text_en", "text_ja", "voice",
    "background", "sprite", "position", "keepCharacters", "next", "ending",
    "inputVariable", "inputPlaceholder_zh", "inputPlaceholder_en",
    "inputButtonText_zh", "inputButtonText_en",
    "secretName", "secretStory", "secretNext",
]
for _i in range(1, MAX_CHOICES + 1):
    STORY_HEADERS += [f"choice{_i}_zh", f"choice{_i}_en", f"choice{_i}_next"]

PET_HEADERS = ["section", "index", "text_zh", "text_en", "text_ja", "voice", "sprite", "duration"]
AVG_HEADERS = ["section", "index", "text_zh", "text_en", "text_ja", "voice", "sprite"]
META_HEADERS = ["file", "field", "value_zh", "value_en", "note"]

HEADER_FILL = PatternFill("solid", fgColor="4F81BD")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


# ----------------------------- 公共工具 -----------------------------

def strip_json_comments(text):
    """去掉 // 与 /* */ 注释，同时正确跳过字符串内容。"""
    out = []
    i, n = 0, len(text)
    in_str = False
    while i < n:
        c = text[i]
        if in_str:
            out.append(c)
            if c == "\\" and i + 1 < n:
                out.append(text[i + 1])
                i += 2
                continue
            if c == '"':
                in_str = False
            i += 1
            continue
        if c == '"':
            in_str = True
            out.append(c)
            i += 1
            continue
        if c == "/" and i + 1 < n and text[i + 1] == "/":
            while i < n and text[i] != "\n":
                i += 1
            continue
        if c == "/" and i + 1 < n and text[i + 1] == "*":
            i += 2
            while i + 1 < n and not (text[i] == "*" and text[i + 1] == "/"):
                i += 1
            i += 2
            continue
        out.append(c)
        i += 1
    return "".join(out)


def load_json(path, ordered=True):
    txt = Path(path).read_text(encoding="utf-8")
    kw = {"object_pairs_hook": OrderedDict} if ordered else {}
    return json.loads(strip_json_comments(txt), **kw)


def save_json(path, data):
    path = Path(path)
    if path.exists():
        # 备份到独立目录，避免 .bak 混进 StreamingAssets 被打包。
        dest = BACKUP_DIR / _BACKUP_STAMP / path.name
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(path, dest)
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def s(v):
    """单元格值 -> 去空白字符串。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def num(v):
    """单元格值 -> int / None。"""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def style_sheet(ws, headers, wrap_cols):
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    widths = {"text_zh": 46, "text_en": 46, "text_ja": 46}
    for col, name in enumerate(headers, start=1):
        letter = get_column_letter(col)
        if name.startswith("text") or name.startswith("choice") or name == "value_zh" or name == "value_en":
            ws.column_dimensions[letter].width = widths.get(name, 30)
        elif name in ("id", "secretStory", "secretNext", "note"):
            ws.column_dimensions[letter].width = 18
        else:
            ws.column_dimensions[letter].width = 14
    ws._wrap_cols = wrap_cols


def write_row(ws, headers, values, row_idx, wrap_cols):
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=col, value=values.get(name, ""))
        if name in wrap_cols:
            cell.alignment = WRAP


# ----------------------------- 导出 JSON -> Excel -----------------------------

def node_sprite(node):
    chars = node.get("characters") or []
    if chars:
        return s(chars[0].get("sprite")), s(chars[0].get("position") or "center")
    return "", ""


def export_story(wb, sheet_name, zh_file, en_file, ja):
    zh = load_json(SA / zh_file)
    en_map = {}
    en_path = SA / en_file
    if en_path.exists():
        en = load_json(en_path)
        en_map = {n.get("id"): n for n in en.get("nodes", [])}

    ws = wb.create_sheet(sheet_name)
    wrap_cols = {"text_zh", "text_en", "text_ja"} | {
        f"choice{i}_zh" for i in range(1, MAX_CHOICES + 1)
    } | {f"choice{i}_en" for i in range(1, MAX_CHOICES + 1)}
    style_sheet(ws, STORY_HEADERS, wrap_cols)

    row_idx = 2
    for node in zh.get("nodes", []):
        nid = node.get("id", "")
        en_node = en_map.get(nid, {})
        sprite, position = node_sprite(node)
        vals = {
            "id": nid,
            "speaker_zh": node.get("speaker", ""),
            "speaker_en": en_node.get("speaker", ""),
            "text_zh": node.get("text", ""),
            "text_en": en_node.get("text", ""),
            "text_ja": ja.get(nid, ""),
            "voice": node.get("voice", ""),
            "background": node.get("background", ""),
            "sprite": sprite,
            "position": position,
            "keepCharacters": "TRUE" if node.get("keepCharacters") else "",
            "next": node.get("next", ""),
            "ending": node.get("ending", ""),
            "inputVariable": node.get("inputVariable", ""),
            "inputPlaceholder_zh": node.get("inputPlaceholder", "") if node.get("inputVariable") else "",
            "inputPlaceholder_en": en_node.get("inputPlaceholder", "") if node.get("inputVariable") else "",
            "inputButtonText_zh": node.get("inputButtonText", "") if node.get("inputVariable") else "",
            "inputButtonText_en": en_node.get("inputButtonText", "") if node.get("inputVariable") else "",
            "secretName": node.get("secretName", ""),
            "secretStory": node.get("secretStory", ""),
            "secretNext": node.get("secretNext", ""),
        }
        zh_choices = node.get("choices") or []
        en_choices = en_node.get("choices") or []
        for i in range(MAX_CHOICES):
            if i < len(zh_choices):
                vals[f"choice{i+1}_zh"] = zh_choices[i].get("text", "")
                vals[f"choice{i+1}_next"] = zh_choices[i].get("next", "")
            if i < len(en_choices):
                vals[f"choice{i+1}_en"] = en_choices[i].get("text", "")
        write_row(ws, STORY_HEADERS, vals, row_idx, wrap_cols)
        row_idx += 1


def export_pet(wb, sheet_name, zh_file, en_file, ja):
    zh = load_json(SA / zh_file)
    en = load_json(SA / en_file) if (SA / en_file).exists() else {}
    ws = wb.create_sheet(sheet_name)
    wrap_cols = {"text_zh", "text_en", "text_ja"}
    style_sheet(ws, PET_HEADERS, wrap_cols)
    row_idx = 2
    for section in PET_SECTIONS:
        zh_list = zh.get(section) or []
        en_list = en.get(section) or []
        for i, item in enumerate(zh_list):
            en_item = en_list[i] if i < len(en_list) else {}
            voice = item.get("voice", "")
            vals = {
                "section": section,
                "index": i + 1,
                "text_zh": item.get("text", ""),
                "text_en": en_item.get("text", ""),
                "text_ja": ja.get(voice, ""),
                "voice": voice,
                "sprite": item.get("sprite", ""),
                "duration": item.get("duration", ""),
            }
            write_row(ws, PET_HEADERS, vals, row_idx, wrap_cols)
            row_idx += 1


def export_avg(wb, ja):
    zh = load_json(SA / AVG_FILES[0])
    en = load_json(SA / AVG_FILES[1]) if (SA / AVG_FILES[1]).exists() else {}
    ws = wb.create_sheet("ending_avg")
    wrap_cols = {"text_zh", "text_en", "text_ja"}
    style_sheet(ws, AVG_HEADERS, wrap_cols)
    row_idx = 2
    for section in AVG_SECTIONS:
        zh_list = zh.get(section) or []
        en_list = en.get(section) or []
        for i, item in enumerate(zh_list):
            en_item = en_list[i] if i < len(en_list) else {}
            voice = item.get("voice", "")
            vals = {
                "section": section,
                "index": i + 1,
                "text_zh": item.get("text", ""),
                "text_en": en_item.get("text", ""),
                "text_ja": ja.get(voice, ""),
                "voice": voice,
                "sprite": item.get("sprite", ""),
            }
            write_row(ws, AVG_HEADERS, vals, row_idx, wrap_cols)
            row_idx += 1


def export_meta(wb):
    ws = wb.create_sheet("meta")
    style_sheet(ws, META_HEADERS, {"value_zh", "value_en"})
    rows = []
    for key, (zh_file, en_file) in STORY_FILES.items():
        zh = load_json(SA / zh_file)
        en = load_json(SA / en_file) if (SA / en_file).exists() else {}
        rows.append({"file": zh_file, "field": "startNode",
                     "value_zh": zh.get("startNode", ""), "value_en": "", "note": "起始节点 id（不翻译）"})
        rows.append({"file": zh_file, "field": "tamperNode",
                     "value_zh": zh.get("tamperNode", ""), "value_en": "", "note": "篡改彩蛋节点 id（不翻译）"})
        rows.append({"file": zh_file, "field": "historyBlockedMessage",
                     "value_zh": zh.get("historyBlockedMessage", ""),
                     "value_en": en.get("historyBlockedMessage", ""), "note": "点历史记录时显示的话（空=正常显示历史）"})
    avg_zh = load_json(SA / AVG_FILES[0])
    avg_en = load_json(SA / AVG_FILES[1]) if (SA / AVG_FILES[1]).exists() else {}
    for field, note in [("speaker", "AVG 说话人名"),
                        ("sleepText", "AVG 睡着时气泡文字"),
                        ("switchButtonText", "AVG 切回桌宠按钮文字")]:
        rows.append({"file": "ending_avg.json", "field": field,
                     "value_zh": avg_zh.get(field, ""), "value_en": avg_en.get(field, ""), "note": note})
    pet_zh = load_json(SA / PET_FILES["pet"][0])
    pet_en = load_json(SA / PET_FILES["pet"][1]) if (SA / PET_FILES["pet"][1]).exists() else {}
    rows.append({"file": "pet_dialogues.json", "field": "sleepText",
                 "value_zh": pet_zh.get("sleepText", ""), "value_en": pet_en.get("sleepText", ""),
                 "note": "桌宠睡着气泡文字（真假两套同用）"})
    for row_idx, vals in enumerate(rows, start=2):
        write_row(ws, META_HEADERS, vals, row_idx, {"value_zh", "value_en"})


def do_export(xlsx_path):
    ja = load_json(JA_PATH, ordered=True) if JA_PATH.exists() else OrderedDict()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (zh_file, en_file) in STORY_FILES.items():
        export_story(wb, name, zh_file, en_file, ja)
    for name, (zh_file, en_file) in PET_FILES.items():
        export_pet(wb, name, zh_file, en_file, ja)
    export_avg(wb, ja)
    export_meta(wb)
    Path(xlsx_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    print(f"已导出表格 -> {xlsx_path}")
    print("在 Excel 里改完文本后，运行:  python tools/text_table/text_table.py import")


# ----------------------------- 导入 Excel -> JSON -----------------------------

def read_sheet(ws):
    headers = [s(c.value) for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2):
        vals = {headers[i]: r[i].value for i in range(len(headers)) if i < len(r)}
        rows.append(vals)
    return rows


ORDER = ["id", "background", "characters", "keepCharacters", "speaker", "text",
         "voice", "inputVariable", "inputPlaceholder", "inputButtonText",
         "secretName", "secretStory", "secretNext", "next", "choices", "ending",
         "setVariables", "conditionalNexts", "editorX", "editorY"]


def ordered_node(d):
    out = OrderedDict()
    for k in ORDER:
        if k in d:
            out[k] = d[k]
    for k, v in d.items():
        if k not in out:
            out[k] = v
    return out


def import_story(ws, zh_file, en_file, ja):
    base = load_json(SA / zh_file)
    base_index = {n.get("id"): n for n in base.get("nodes", [])}
    zh_nodes, en_nodes = [], []
    for row in read_sheet(ws):
        nid = s(row.get("id"))
        if not nid:
            continue
        old = base_index.get(nid, {})
        node = OrderedDict()
        node["id"] = nid
        if s(row.get("background")):
            node["background"] = s(row.get("background"))
        # 立绘
        if s(row.get("keepCharacters")).upper() in ("TRUE", "1", "YES"):
            node["keepCharacters"] = True
        elif s(row.get("sprite")):
            node["characters"] = [OrderedDict([
                ("sprite", s(row.get("sprite"))),
                ("position", s(row.get("position")) or "center"),
            ])]
        else:
            if "characters" in old:
                node["characters"] = old["characters"]
            if old.get("keepCharacters"):
                node["keepCharacters"] = True
        if s(row.get("speaker_zh")):
            node["speaker"] = s(row.get("speaker_zh"))
        if s(row.get("text_zh")):
            node["text"] = s(row.get("text_zh"))
        if s(row.get("voice")):
            node["voice"] = s(row.get("voice"))
        # 输入框
        if s(row.get("inputVariable")):
            node["inputVariable"] = s(row.get("inputVariable"))
            if s(row.get("inputPlaceholder_zh")):
                node["inputPlaceholder"] = s(row.get("inputPlaceholder_zh"))
            if s(row.get("inputButtonText_zh")):
                node["inputButtonText"] = s(row.get("inputButtonText_zh"))
        # 隐藏真名分支
        if s(row.get("secretName")):
            node["secretName"] = s(row.get("secretName"))
            node["secretStory"] = s(row.get("secretStory"))
            node["secretNext"] = s(row.get("secretNext"))
        # 选项 / next
        choices = []
        for i in range(1, MAX_CHOICES + 1):
            ctext = s(row.get(f"choice{i}_zh"))
            cnext = s(row.get(f"choice{i}_next"))
            if ctext or cnext:
                choices.append(OrderedDict([("text", ctext), ("next", cnext)]))
        if choices:
            node["choices"] = choices
        elif s(row.get("next")):
            node["next"] = s(row.get("next"))
        if s(row.get("ending")):
            node["ending"] = s(row.get("ending"))
        # 表格不管理、从旧 JSON 保留的结构字段
        for keep in ("setVariables", "conditionalNexts", "editorX", "editorY"):
            if keep in old:
                node[keep] = old[keep]
        zh_nodes.append(ordered_node(node))

        # 英文节点：结构照抄中文，只换文本
        en_node = OrderedDict(node)
        if s(row.get("speaker_en")):
            en_node["speaker"] = s(row.get("speaker_en"))
        if s(row.get("text_en")):
            en_node["text"] = s(row.get("text_en"))
        if "inputVariable" in node:
            if s(row.get("inputPlaceholder_en")):
                en_node["inputPlaceholder"] = s(row.get("inputPlaceholder_en"))
            if s(row.get("inputButtonText_en")):
                en_node["inputButtonText"] = s(row.get("inputButtonText_en"))
        if "choices" in en_node:
            en_choices = []
            for i in range(1, MAX_CHOICES + 1):
                ctext = s(row.get(f"choice{i}_en")) or s(row.get(f"choice{i}_zh"))
                cnext = s(row.get(f"choice{i}_next"))
                if ctext or cnext:
                    en_choices.append(OrderedDict([("text", ctext), ("next", cnext)]))
            en_node["choices"] = en_choices
        en_nodes.append(ordered_node(en_node))

        # 日语配音（按节点 id，即 wav 文件名）
        if s(row.get("text_ja")):
            ja[nid] = s(row.get("text_ja"))

    zh_out = OrderedDict(base)
    zh_out["nodes"] = zh_nodes
    save_json(SA / zh_file, zh_out)

    en_base = load_json(SA / en_file) if (SA / en_file).exists() else OrderedDict(base)
    en_out = OrderedDict(en_base)
    for meta_key in ("startNode", "tamperNode"):
        en_out[meta_key] = zh_out.get(meta_key, en_out.get(meta_key, ""))
    en_out["nodes"] = en_nodes
    save_json(SA / en_file, en_out)


def import_pet(ws, zh_file, en_file, ja):
    zh = load_json(SA / zh_file)
    en = load_json(SA / en_file) if (SA / en_file).exists() else OrderedDict(zh)
    sections_zh = {sec: [] for sec in PET_SECTIONS}
    sections_en = {sec: [] for sec in PET_SECTIONS}
    for row in read_sheet(ws):
        sec = s(row.get("section"))
        if sec not in PET_SECTIONS:
            continue
        voice = s(row.get("voice"))
        item = OrderedDict()
        item["text"] = s(row.get("text_zh"))
        if voice:
            item["voice"] = voice
        if s(row.get("sprite")):
            item["sprite"] = s(row.get("sprite"))
        d = num(row.get("duration"))
        if d is not None:
            item["duration"] = d
        sections_zh[sec].append(item)

        item_en = OrderedDict(item)
        item_en["text"] = s(row.get("text_en")) or s(row.get("text_zh"))
        sections_en[sec].append(item_en)

        if voice and s(row.get("text_ja")):
            ja[voice] = s(row.get("text_ja"))

    zh_out = OrderedDict(zh)
    en_out = OrderedDict(en)
    for sec in PET_SECTIONS:
        zh_out[sec] = sections_zh[sec]
        en_out[sec] = sections_en[sec]
    save_json(SA / zh_file, zh_out)
    save_json(SA / en_file, en_out)


def import_avg(ws, ja):
    zh_file, en_file = AVG_FILES
    zh = load_json(SA / zh_file)
    en = load_json(SA / en_file) if (SA / en_file).exists() else OrderedDict(zh)
    sec_zh = {sec: [] for sec in AVG_SECTIONS}
    sec_en = {sec: [] for sec in AVG_SECTIONS}
    for row in read_sheet(ws):
        sec = s(row.get("section"))
        if sec not in AVG_SECTIONS:
            continue
        voice = s(row.get("voice"))
        item = OrderedDict()
        item["text"] = s(row.get("text_zh"))
        if voice:
            item["voice"] = voice
        if s(row.get("sprite")):
            item["sprite"] = s(row.get("sprite"))
        sec_zh[sec].append(item)
        item_en = OrderedDict(item)
        item_en["text"] = s(row.get("text_en")) or s(row.get("text_zh"))
        sec_en[sec].append(item_en)
        if voice and s(row.get("text_ja")):
            ja[voice] = s(row.get("text_ja"))
    zh_out = OrderedDict(zh)
    en_out = OrderedDict(en)
    for sec in AVG_SECTIONS:
        zh_out[sec] = sec_zh[sec]
        en_out[sec] = sec_en[sec]
    save_json(SA / zh_file, zh_out)
    save_json(SA / en_file, en_out)


def apply_meta(ws):
    """meta 页签里少量固定文本 / 节点 id 的覆盖。在其它导入之后执行。"""
    cache = {}

    def get(path):
        if path not in cache:
            cache[path] = load_json(SA / path)
        return cache[path]

    def put(d, field, val):
        # 空值不写空键：有值则设，无值则删（这样也能用来清空某个字段）。
        if val:
            d[field] = val
        else:
            d.pop(field, None)

    for row in read_sheet(ws):
        f = s(row.get("file"))
        field = s(row.get("field"))
        vz = s(row.get("value_zh"))
        ve = s(row.get("value_en"))
        if not f or not field:
            continue
        if f in ("story.json", "story_true.json"):
            en_file = f.replace(".json", "_en.json")
            zh = get(f)
            en = get(en_file)
            if field in ("startNode", "tamperNode"):
                put(zh, field, vz)
                put(en, field, vz)
            elif field == "historyBlockedMessage":
                put(zh, field, vz)
                put(en, field, ve or vz)
        elif f == "ending_avg.json":
            zh = get("ending_avg.json")
            en = get("ending_avg_en.json")
            put(zh, field, vz)
            put(en, field, ve or vz)
        elif f == "pet_dialogues.json" and field == "sleepText":
            for pf in ("pet_dialogues.json", "pet_dialogues_en.json",
                       "pet_dialogues_true.json", "pet_dialogues_true_en.json"):
                if (SA / pf).exists():
                    cfg = get(pf)
                    put(cfg, "sleepText", (ve or vz) if pf.endswith("_en.json") else vz)

    for path, data in cache.items():
        save_json(SA / path, data)


def do_import(xlsx_path):
    if not Path(xlsx_path).exists():
        print(f"找不到表格 {xlsx_path}，请先运行 export 生成。")
        sys.exit(1)
    wb = openpyxl.load_workbook(xlsx_path)
    ja = load_json(JA_PATH, ordered=True) if JA_PATH.exists() else OrderedDict()

    for name, (zh_file, en_file) in STORY_FILES.items():
        if name in wb.sheetnames:
            import_story(wb[name], zh_file, en_file, ja)
    for name, (zh_file, en_file) in PET_FILES.items():
        if name in wb.sheetnames:
            import_pet(wb[name], zh_file, en_file, ja)
    if "ending_avg" in wb.sheetnames:
        import_avg(wb["ending_avg"], ja)
    if "meta" in wb.sheetnames:
        apply_meta(wb["meta"])

    save_json(JA_PATH, ja)
    print("已写回 JSON:")
    print("  Assets/StreamingAssets/*.json  和  tools/tts/voice_lines_ja.json")
    print(f"旧文件备份在: tools/text_table/backups/{_BACKUP_STAMP}/")
    print("如需重新生成配音，运行:  python tools/tts/synthesize.py")


def main():
    ap = argparse.ArgumentParser(description="游戏文本表格 <-> JSON 双向转换")
    ap.add_argument("action", choices=["export", "import"])
    ap.add_argument("--xlsx", default=str(DEFAULT_XLSX))
    args = ap.parse_args()
    if args.action == "export":
        do_export(args.xlsx)
    else:
        do_import(args.xlsx)


if __name__ == "__main__":
    main()
